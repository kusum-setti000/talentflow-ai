"""Local document parsing. Never sends raw files to an external service."""
from io import BytesIO
from pathlib import Path
import zipfile
import pymupdf
from docx import Document
from talentflow.config import Settings
from talentflow.domain.errors import InputError


def read_document(filename: str, content: bytes, settings: Settings) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise InputError("Only PDF and DOCX documents are supported.")
    if not content or len(content) > settings.max_file_mb * 1024 * 1024:
        raise InputError(f"File must be non-empty and at most {settings.max_file_mb} MB.")
    try:
        if suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise InputError("The uploaded file is not a valid PDF.")
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.needs_pass or not 1 <= len(pdf) <= settings.max_pages:
                    raise InputError("Encrypted or oversized PDFs are not supported.")
                text = "\n".join(page.get_text() for page in pdf)
        else:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(m.file_size for m in members) > 40_000_000:
                    raise InputError("DOCX expanded content is too large.")
                if any(m.flag_bits & 1 for m in members) or "word/document.xml" not in archive.namelist():
                    raise InputError("Encrypted or invalid DOCX file.")
            document = Document(BytesIO(content))
            text = "\n".join([*(p.text for p in document.paragraphs), *(cell.text for table in document.tables for row in table.rows for cell in row.cells)])
    except InputError:
        raise
    except Exception as exc:
        raise InputError("Unable to read this document. Upload a valid PDF/DOCX or paste its text.") from exc
    text = text.strip()
    if len(text) > settings.max_characters:
        raise InputError("Extracted document exceeds the text limit.")
    if not text:
        raise InputError("No extractable text found. Paste the text or upload a text-based PDF/DOCX.")
    return text
